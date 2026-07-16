"""
Text extraction from resume files (PDF, DOCX)

Primary: pdfplumber (PDF), python-docx (DOCX)
Fallback: PyMuPDF (fitz) for PDFs where pdfplumber returns suspiciously
little text – common with resumes exported from design tools (Canva,
Figma) that flatten text into odd encodings pdfplumber sometimes mangles
"""
from __future__ import annotations

import os
import re

import fitz  # PyMuPDF
import pdfplumber
from docx import Document


def extract_from_pdf(path: str) -> str:
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    text = "\n".join(text_parts).strip()

    # suspiciously short extraction (e.g. a design-tool PDF) – retry with
    # PyMuPDF before giving up.
    if len(text) < 50:
        text = _extract_from_pdf_pymupdf(path)

    return text


def _extract_from_pdf_pymupdf(path: str) -> str:
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def extract_from_docx(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    # table cells are common in resume templates (skills grids, two-column
    # contact blocks) and get silently dropped if you only read paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(p for p in parts if p.strip())


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(path)
    elif ext == ".docx":
        return extract_from_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext} (expected .pdf or .docx)")


def clean_text(raw_text: str) -> str:
    """Light normalization before NER: collapse repeated whitespace, strip
    stray control chars pdfplumber sometimes leaves behind. Line breaks are
    kept – both NER and the years-of-experience regex benefit from
    structure (bullet points, section breaks)"""
    text = raw_text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
